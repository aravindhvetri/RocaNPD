"""Generate the ROCA NPD end-to-end code review Word document."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

TEAL = RGBColor(0x1D, 0x4E, 0x56)
ACCENT = RGBColor(0x40, 0x91, 0x9D)
DARK = RGBColor(0x16, 0x2C, 0x34)
MUTED = RGBColor(0x4B, 0x55, 0x63)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEADER_BG = "1D4E56"
ROW_ALT = "E8F4F6"
OUTPUT = Path(__file__).with_name("ROCA_NPD_End_to_End_Code_Review.docx")


def set_run(run, *, size=11, bold=False, color=DARK, font="Calibri"):
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), font)


def shade_cell(cell, hex_color: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_text(cell, text: str, *, bold=False, color=DARK, size=10, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, color=WHITE, size=10)
        shade_cell(table.rows[0].cells[i], HEADER_BG)
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            set_cell_text(table.rows[r_idx + 1].cells[c_idx], value, size=10)
            if r_idx % 2 == 1:
                shade_cell(table.rows[r_idx + 1].cells[c_idx], ROW_ALT)
    doc.add_paragraph()


def heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = TEAL if level > 1 else DARK


def para(doc: Document, text: str, *, bold=False) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    set_run(run, bold=bold)


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        for run in p.runs:
            set_run(run, size=11)


def code(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Cm(0.5)
    run = p.add_run(text)
    set_run(run, size=9, font="Consolas", color=DARK)


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("ROCA NPD")
    set_run(run, size=28, bold=True, color=TEAL, font="Calibri")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("End-to-End Code Review Guide")
    set_run(run, size=18, bold=True, color=ACCENT)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(
        "New Product Development Request Tracker  |  18 September 2026  |  SPFx + React + Redux + SharePoint"
    )
    set_run(run, size=11, color=MUTED)

    para(
        doc,
        "This document explains the complete ROCA NPD application for code review: "
        "what each component is for, how it works, where data is fetched, which SharePoint "
        "lists own that data, how components connect through Redux and services, and the "
        "full request flow from login to approval. It describes the current codebase as implemented, "
        "including items still pending (Material Group screens, SAP posting, Reports).",
    )

    heading(doc, "1. What this application is", 1)
    para(
        doc,
        "ROCA NPD is a SharePoint Framework (SPFx) client-side web part named "
        "“NPD REQUESTS — New Product Development”. Users create New Product Development requests, "
        "save drafts, submit them into a sequential approval chain (Initiator → Vertical Head → MIS Coordinator), "
        "and review requests on role-scoped list screens. Administration of Lookup Type, Lookup, "
        "Workflow Configuration, and Brand Material Extension lives in the same app. Plant, Role, "
        "and Approver master data are not maintained here — they are read from the separate ROCA site.",
    )
    bullets(
        doc,
        [
            "NPD request create / draft / submit / view / approve / rework / reject.",
            "Role-based navigation and data scope from ROCA ApproversMaster (union of roles).",
            "Admin masters for lookups, workflow steps, and brand-material extension.",
            "Email notification on submit and Vertical Head actions (Graph CID logo).",
            "Material Group, SAP Post, Material Master auto-populate, and Reports are designed but not implemented as full screens yet.",
        ],
    )

    heading(doc, "2. Technology stack", 1)
    add_table(
        doc,
        ["Layer", "Choice", "Where it lives"],
        [
            ["Host", "SharePoint Online + SPFx 1.23.2", "src/webparts/rocaNpd/RocaNpdWebPart.ts"],
            ["UI", "React 17 + PrimeReact 10, wrapped", "components/common/controls/"],
            ["State", "Redux Toolkit", "src/store/"],
            ["SharePoint API", "PnPjs v4 (spfi + SPFx)", "External/CommonServices/SPServices.ts"],
            ["Routing", "react-router-dom HashRouter", "#/npd/..., #/admin/..."],
            ["Styling", "SCSS modules + theme.scss", "External/CommonServices/theme.scss"],
            ["Standards", "Mandatory MD docs", "src/Documents/"],
        ],
    )
    para(
        doc,
        "Feature screens must not import primereact/* directly. They use the common wrappers. "
        "Components must not call SharePoint; they dispatch thunks, thunks call services, "
        "services call SPServices.",
    )

    heading(doc, "3. How the application starts", 1)
    para(doc, "Bootstrap is a straight line. Nothing in a screen should re-initialize PnP.")
    code(
        doc,
        "RocaNpdWebPart.onInit\n"
        "  → loadApplicationStyles (PrimeReact + Poppins + theme)\n"
        "RocaNpdWebPart.render\n"
        "  → React.createElement(RocaNpd, { context })\n"
        "RocaNpd constructor\n"
        "  → spfi().using(SPFx(context))\n"
        "  → setupSP(sp) and setupSpfxContext(context)\n"
        "  → <Provider store> → MainComponent\n"
        "MainComponent\n"
        "  → dispatch(initializeApp(spfxContext))\n"
        "  → HashRouter → AppShell (Header + SideNavigation) → AppRoutes",
    )
    para(
        doc,
        "initializeApp (src/store/thunks/appThunks.ts) writes the logged-in user into appSlice "
        "(display name, email, login name, userId, site URL), then calls resolveUserAccess "
        "in roleService.ts. That function:",
    )
    bullets(
        doc,
        [
            "Adds Admin if the user is in the NPD-site SharePoint group “Admins”.",
            "Loads ROCA ApproversMaster (System, Role, Brand, Users expanded).",
            "Matches the current user by email only (normalizeEmail). It does not use NPD-site user IDs or WorkflowJSON.",
            "Collects every matching role as a union (Initiator, Vertical Head, MIS Coordinator, Consultant) plus per-role brand lists.",
            "Stores the result in appSlice.assignedRoles / mappedBrands / npdInitiatorBrands / npdVerticalHeadBrands, etc.",
        ],
    )
    para(
        doc,
        "Until that finishes, MainComponent shows LoaderOverlay. ProtectedRoute then blocks any "
        "route the resolved roles cannot access. Side navigation is filtered by the same access object "
        "(useFilteredNavigation → permissionService.filterNavigationByRoles).",
    )

    heading(doc, "4. Layered architecture", 1)
    code(
        doc,
        "Screen / layout component\n"
        "        ↓  dispatch / useAppSelector\n"
        "Redux thunk (store/thunks)\n"
        "        ↓  calls\n"
        "Service (External/CommonServices)\n"
        "        ↓  PnP via SPServices, sometimes Web([sp.web, rocaSiteUrl])\n"
        "SharePoint list item\n"
        "        ↓  mapped DTO\n"
        "Slice extraReducer writes state\n"
        "        ↓\n"
        "Screen re-renders",
    )
    para(
        doc,
        "This is the only allowed data path (ProjectStandards ADR-005). If you find PnP inside a "
        "React component during review, that is a defect.",
    )

    heading(doc, "5. Folder map (what lives where)", 1)
    add_table(
        doc,
        ["Path", "Responsibility"],
        [
            ["src/webparts/rocaNpd/RocaNpdWebPart.ts", "SPFx lifecycle, styles, render React root"],
            ["src/webparts/rocaNpd/components/RocaNpd.tsx", "PnP setup + Redux Provider"],
            [".../components/MainComponent.tsx", "initializeApp, HashRouter, AppShell, routes"],
            [".../components/layout/", "Header, AppShell, SideNavigation, RoleIndicator"],
            [".../components/routes/", "AppRoutes, ProtectedRoute, Unauthorized"],
            [".../components/npd/newRequest/", "NPD form (create / view / edit / approve)"],
            [".../components/npd/requestList/", "All + Approved dashboards + Workflow Status"],
            [".../components/npd/draftRework/", "Draft/Rework and Pending list screens"],
            [".../components/admin/", "Lookup Type, Lookup, Workflow Config, Brand Extension"],
            [".../components/common/controls/", "PrimeReact wrappers used by every screen"],
            ["src/store/", "Redux store, slices, typed hooks, thunks"],
            ["src/External/CommonServices/", "Config, PnP, all SharePoint/business services"],
            ["src/Documents/", "TechnicalArchitecture, ProjectStandards, TaskList, Checklist"],
        ],
    )

    heading(doc, "6. Redux store — how state is connected", 1)
    para(
        doc,
        "configureStore in src/store/index.ts registers five slices. Screens select with "
        "useAppSelector and dispatch with useAppDispatch (src/store/hooks.ts).",
    )
    add_table(
        doc,
        ["Slice", "Holds", "Filled by"],
        [
            ["appSlice", "User identity, roles, brands, init flag", "initializeApp / roleService"],
            ["uiSlice", "Sidebar collapsed, toasts-related UI flags", "layout actions"],
            ["adminSlice", "Lookup Type, Lookup, Workflow Config, Brand Extension rows", "admin thunks"],
            ["npdFormSlice", "Open request: general info, item rows, lookup options, submit progress", "npdFormThunks"],
            ["npdRequestSlice", "List rows for Draft, Pending, All/Approved dashboards", "npdRequestThunks"],
        ],
    )
    para(
        doc,
        "npdFormSlice is the working copy of one request. npdRequestSlice is the list cache. "
        "They are separate so opening a form does not wipe the dashboard, and saving a draft "
        "does not require the list to own form field state.",
    )

    heading(doc, "7. Where the data comes from", 1)
    para(
        doc,
        "There are two SharePoint sites. The NPD site hosts this web part and the NPD lists. "
        "The ROCA master site hosts shared masters. Cross-site reads use requireRocaMasterSiteUrl() "
        "from the SPFx site URL (not window.origin alone) and PnP Web([sp.web, rocaSiteUrl]).",
    )
    heading(doc, "7.1 NPD site lists", 2)
    add_table(
        doc,
        ["List (Config.ListNames)", "Used for", "Primary service"],
        [
            ["NPD_Request", "Request header: Brand, MaterialType, Plant, Status, WorkFlowJSON, Title/Request ID", "npdRequestGeneralInfoService.ts"],
            ["NPD_ItemDetails", "Line items; lookup to NPD_Request via RequestGeneralInfoId", "npdItemDetailsService.ts"],
            ["NPD_ApproverComments", "Rework/Reject comments + acting user Person field", "npdApproverCommentsService.ts"],
            ["NPD_LookupType", "Admin Lookup Type master", "lookupTypeService.ts"],
            ["NPD_Lookup", "Admin Lookup values AND Item Details MultiSelect options", "lookupService.ts + lookupOptionUtils.ts"],
            ["NPD_WorkflowConfig", "Approval chain: RequestType, CurrentRole, NextRole, order", "workflowConfigurationService.ts"],
            ["NPD_BrandMaterialExtensionMaster", "Brand → plant/warehouse extension mapping", "brandMaterialExtensionService.ts"],
            ["NPD_Templates library", "Excel import templates + approval email HTML template", "templateService.ts"],
        ],
    )
    heading(doc, "7.2 ROCA master site lists", 2)
    add_table(
        doc,
        ["List", "Used for", "Primary service"],
        [
            ["ApproversMaster", "Who has which role on which brand/system; initiator Brand dropdown", "roleService.ts, npdApproverAssignmentService.ts, npdFormDataService.ts"],
            ["PlantMaster", "Finished Products Plant/Source (PlantType=Factory, active, PlantCode)", "npdFormDataService.ts / rocaMasterDataService.ts"],
            ["RoleMaster", "Workflow Config Next Role options (NPD excludes Consultant)", "workflowConfigurationService.ts"],
            ["BrandMaster", "Brand titles when needed as ROCA lookup source", "rocaMasterDataService.ts"],
        ],
    )
    para(
        doc,
        "ApproversMaster matching is email-only after expand (System/Title, Role/Title, Brand/Title, Users/EMail). "
        "Do not filter those lookup/person fields in OData — load and filter in TypeScript "
        "(R-NPD05a). Vertical Head brands are brand-scoped. MIS Coordinator is not brand-scoped for pending work.",
    )

    heading(doc, "8. Service catalog (review this layer first)", 1)
    para(
        doc,
        "Every SharePoint call funnels through SPServices.ts (add/update/delete/read/batch, getAllUsers, ensureUser, groups). "
        "Business services sit on top of it.",
    )
    add_table(
        doc,
        ["Service file", "What it does"],
        [
            ["Config.ts", "List names, field names, labels, roles, routes, statuses — never hardcode these in UI"],
            ["Interface.ts", "DTOs: INpdRequestGeneralInfo, INpdWorkflowStepJson, IResolvedUserAccess, etc."],
            ["roleService.ts", "Resolve assigned roles + brand maps from ApproversMaster + Admin group"],
            ["permissionService.ts", "hasRole, canViewRequest, canActOnPendingNpdStep, filterNavigationByRoles, ProtectedRoute helpers"],
            ["personFieldUtils.ts", "normalizeEmail, extractPersonEmails, approversRowMatchesEmail"],
            ["npdWorkflowJsonService.ts", "Build, parse, submit, approve/rework/reject WorkFlowJSON; pending role helpers"],
            ["npdApproverAssignmentService.ts", "Emails for a workflow role + brand from ApproversMaster"],
            ["npdRequestGeneralInfoService.ts", "CRUD header; Save Draft; Submit; applyNpdWorkflowAction; Author expand"],
            ["npdRequestListService.ts", "Dashboard / pending / draft-rework queries + item-line counts"],
            ["npdItemDetailsService.ts", "Batch insert/update/delete line items"],
            ["npdItemDetailsImportService.ts", "Excel parse + lookup option validation for Item Details"],
            ["npdRequestIdService.ts", "NPD-YYYY-### generation, skip IsDeleted"],
            ["npdFormDataService.ts", "Initiator brands + factory plants (ROCA)"],
            ["npdNotificationService.ts", "Send approval emails after submit / VH action"],
            ["npdWorkflowStatusView.ts", "Read-only Workflow Status rows + display-name resolution"],
            ["lookupService.ts / lookupTypeService.ts", "Admin CRUD + active lookups for MultiSelect options"],
            ["importService.ts / exportService.ts", "Shared Excel import/export used by Admin and NPD"],
            ["softDelete.ts", "IsDeleted payloads; all deletes are soft"],
            ["rocaSiteUrlResolver.ts", "Map current site URL → ROCA master site URL"],
        ],
    )

    heading(doc, "9. Common UI components (used everywhere)", 1)
    para(
        doc,
        "These wrappers are the only PrimeReact surface. Feature screens import from "
        "components/common/controls. If a pattern appears twice, it belongs here (R-C02).",
    )
    add_table(
        doc,
        ["Component", "Used for"],
        [
            ["Button", "Primary/secondary/text/icon actions (Edit, View, Workflow, Save, Submit)"],
            ["InputText / InputNumber / InputTextarea", "Form fields; NPD Form sets autoComplete=off"],
            ["Dropdown / MultiSelect / ComboBox", "Single and multi lookups; overlay styles from theme"],
            ["DatePicker", "Date fields in admin/filter patterns"],
            ["DataTable", "Every master and request list, plus Item Details grid"],
            ["Dialog / ConfirmDialog", "Forms, Workflow Status, comments, confirmations"],
            ["DeleteConfirmDialog / DeleteBlockedDialog", "Admin soft-delete and dependency-blocked delete"],
            ["Toast", "One-at-a-time validation and success/error messages"],
            ["Tag", "Status chips (Draft, Pending, Rework, Approved, Rejected)"],
            ["LoaderOverlay", "App init, list load, submit line-item progress bar"],
            ["FileUpload", "Excel import"],
            ["ControlField", "Label + required + control layout"],
            ["MasterTablePanel / MasterToolbarSearch", "Admin and NPD list chrome"],
            ["ImportDialog / ImportValidationDialog", "Shared import UX; NPD validation shows S.NO + Error only"],
        ],
    )

    heading(doc, "10. Layout, navigation, and routing", 1)
    add_table(
        doc,
        ["Component", "Role"],
        [
            ["AppShell", "Header + collapsible sidebar + content slot + version label"],
            ["Header", "App title “New Product Development” and user context"],
            ["SideNavigation / Compact variants", "NAV_SECTIONS from navigationConfig.ts, filtered by roles"],
            ["RoleIndicator", "Shows resolved roles (union, not a switch that overrides others)"],
            ["useFilteredNavigation.ts", "Applies permissionService.filterNavigationByRoles(access)"],
            ["AppRoutes.tsx", "Maps hash paths to implemented screens or RoutePlaceholder"],
            ["ProtectedRoute.tsx", "Blocks URL access if the resolved role cannot open that route"],
            ["Unauthorized.tsx", "Fallback when the user has no access"],
        ],
    )
    para(doc, "Implemented hash routes:")
    add_table(
        doc,
        ["Route", "Screen component", "Who typically sees it"],
        [
            ["/npd/new", "NpdRequestForm", "Initiator create/edit; all roles view; VH/MIS edit when acting"],
            ["/npd/all", "NpdRequestDashboard variant=all", "Initiator, VH, MIS, Admin"],
            ["/npd/pending", "NpdDraftRework (pending fetch)", "Initiator (own), VH (brand pending), MIS (pending role)"],
            ["/npd/approved", "NpdRequestDashboard variant=approved", "Same visibility as All, filtered to Approved"],
            ["/npd/draft-rework", "NpdDraftRework", "Initiator only (own Draft/Rework)"],
            ["/admin/lookup-type", "LookupTypeMaster", "Admin"],
            ["/admin/lookup", "LookupMaster", "Admin"],
            ["/admin/workflow-config", "WorkflowConfigurationMaster", "Admin"],
            ["/admin/brand-extension", "BrandMaterialExtensionMaster", "Admin"],
            ["MG + Reports + Material Master", "RoutePlaceholder", "Not built yet"],
        ],
    )
    para(
        doc,
        "Vertical Head navigation is strictly All Requests, Pending Approval, Approved Requests. "
        "No New, Draft, Material Group, Admin, or Reports. That filter is role+module based from "
        "ApproversMaster, not from WorkflowJSON or the last selected role.",
    )

    heading(doc, "11. Administration module — component by component", 1)
    para(
        doc,
        "Each admin screen follows the same pattern: Master page → Toolbar (search/add/import/export) "
        "→ Table → Form dialog. Thunks live in lookupTypeThunks, lookupThunks, "
        "workflowConfigurationThunks, brandMaterialExtensionThunks. Soft delete uses IsDeleted. "
        "Delete is blocked when dependencyValidationService finds child records.",
    )
    add_table(
        doc,
        ["Screen", "Components", "Data"],
        [
            ["Lookup Type", "LookupTypeMaster, Toolbar, Table, FormDialog", "NPD_LookupType via lookupTypeService"],
            ["Lookup", "LookupMaster, Toolbar, Table, FormDialog", "NPD_Lookup; type from Lookup Type; Title is the option label used later on the NPD form"],
            ["Workflow Configuration", "WorkflowConfigurationMaster, Table, WorkflowConfigFormDialog", "NPD_WorkflowConfig + RoleMaster for Next Role; builds the NPD approval sequence"],
            ["Brand Material Extension", "BrandMaterialExtensionMaster, Table, FormDialog", "NPD_BrandMaterialExtensionMaster; ROCA brand/plant options"],
        ],
    )
    para(
        doc,
        "Workflow Configuration is what later becomes WorkFlowJSON on a request. "
        "If NPD Request steps are missing, Save Draft / Submit throw “NPD workflow is not configured.”",
    )

    heading(doc, "12. NPD list screens — how they fetch and connect", 1)
    para(
        doc,
        "All four request lists share one table: NpdRequestListTable. Draft/Rework and Pending "
        "wrap it with NpdDraftReworkTable. All and Approved use NpdRequestDashboard.",
    )
    add_table(
        doc,
        ["Screen", "Thunk", "Service filter"],
        [
            ["All Requests", "fetchNpdDashboardList", "fetchNpdDashboardItems → canViewRequest (brand/role union; MIS sees all)"],
            ["Approved Requests", "same thunk, client filter Approved", "Same scope as All, Status = Approved"],
            ["Pending Approval", "fetchNpdPendingList", "Status=Pending AND (user can act on pending role OR initiator owns the row)"],
            ["Draft / ReWork", "fetchNpdDraftReworkList", "Status Draft/Rework AND AuthorEmail = current user"],
        ],
    )
    para(
        doc,
        "List rows come from NPD_Request (Author expanded) plus a summary of NPD_ItemDetails "
        "(product/line count). Columns: Request ID, Brand (MG1), Material Type, Plant, Products, "
        "Status, Created Date, Actions.",
    )
    heading(doc, "12.1 Actions column", 2)
    bullets(
        doc,
        [
            "Edit (pencil) — only if canEditNpdListRow: own Draft/Rework, or current pending VH/MIS step. Opens /npd/new?id=&mode=edit&from=...",
            "Workflow (sitemap) — always. Opens read-only Workflow Status dialog. Does not change data.",
            "View (eye) — always. Opens /npd/new?id=&mode=view&from=... Read-only; never shows Approve/Reject/Rework.",
            "No Delete on request tables.",
        ],
    )
    para(
        doc,
        "from is a short key (all / pending / approved / draft-rework). Cancel/Back on the form "
        "uses resolveNpdFormCancelRoute so the user returns to the list they came from. "
        "Email links use from=pending.",
    )

    heading(doc, "13. NPD form — components and data fetch", 1)
    para(
        doc,
        "Route /npd/new is one form used for New, Draft edit, View, and approver Edit. "
        "useNpdRequestFormController.ts owns load, save, submit, approve, cancel, and footer mode. "
        "NpdRequestForm.tsx renders the shell (form autoComplete=off). Sections are split to stay small.",
    )
    add_table(
        doc,
        ["Component", "What it does", "Data"],
        [
            ["NpdRequestForm", "Page shell, toasts, loader, wires controller", "npdFormSlice"],
            ["NpdGeneralInfoSection", "Brand, Material Type, Plant/Source", "Brands from ApproversMaster; Material Type from Config.NpdMaterialTypes; Plant from PlantMaster or Config.NpdTradedPlantSources"],
            ["NpdItemDetailsSection", "Line-item grid + Import + add/delete row", "NPD_ItemDetails + NPD_Lookup options"],
            ["NpdItemDetailsTableCell / useNpdItemDetailsColumns", "Cell editors (Input/MultiSelect) by column config", "npdItemDetailsConfig.ts"],
            ["NpdItemDetailsImport", "Excel import using shared ImportDialog", "NPD_Templates TemplateType=NPDItemDetails; lookup values must exist"],
            ["NpdFormSectionPanel", "Teal section header chrome", "Presentational"],
            ["NpdRequestFormFooter", "Cancel/Back, Save Draft, Submit, Approve/Rework/Reject/Post to SAP", "Footer mode from role + request status + view/edit"],
            ["NpdApproverCommentDialog", "Comment required for Rework/Reject", "Written to NPD_ApproverComments"],
        ],
    )
    heading(doc, "13.1 Hydration (opening an existing request)", 2)
    code(
        doc,
        "URL: /npd/new?id={Id}&mode=view|edit&from=pending\n"
        "hydrateNpdRequestForm thunk\n"
        "  → npdRequestGeneralInfoService.get by Id\n"
        "  → npdItemDetailsService.fetch by RequestGeneralInfoId\n"
        "  → npdFormSlice: generalInfo, items, requestStatus, WorkFlowJSON steps, title\n"
        "Also in parallel as needed:\n"
        "  fetchNpdInitiatorBrandOptions / fetchNpdPlantSourceOptions / fetchNpdLookupOptions",
    )
    para(
        doc,
        "Do not use setNpdMaterialType when hydrating a saved Plant/Source — that action clears plant. "
        "The controller sets general info fields so an existing plant is preserved (R-NPD10).",
    )
    heading(doc, "13.2 Save Draft", 2)
    para(
        doc,
        "saveNpdDraft → npdRequestGeneralInfoService. Creates or updates NPD_Request. "
        "Status stays Draft on create; existing Rework is preserved. Title stores Brand until a Request ID exists. "
        "WorkFlowJSON is rebuilt every save from NPD_WorkflowConfig + ApproversMaster "
        "(Initiator = current user Status \"\"; approvers Status \"\"). "
        "Item details are batch written. No Request ID yet. No email.",
    )
    heading(doc, "13.3 Submit Request", 2)
    para(
        doc,
        "submitNpdRequest validates (first error toast only). Then:",
    )
    bullets(
        doc,
        [
            "Generate Title = NPD-YYYY-### (npdRequestIdService; skip IsDeleted).",
            "Status = Pending.",
            "applySubmitWorkflowStatuses: first non-Initiator role (usually Vertical Head) → Pending; later roles stay empty; Initiator Status remains \"\".",
            "Patch item RequestGeneralInfoId to the header.",
            "Loader shows “{n} / {m} Line Items Added”.",
            "Email the pending Vertical Head users (Approve / Re-work / Reject buttons in mail). MIS mail does not get those buttons.",
        ],
    )

    heading(doc, "14. Workflow JSON and the Workflow Status dialog", 1)
    para(
        doc,
        "NPD_Request.WorkFlowJSON is an array of { Role, UserEmail, Status }. "
        "It is the live approval tracker. The list already has it parsed as WorkflowSteps on each row.",
    )
    heading(doc, "14.1 How status updates as approvers act", 2)
    code(
        doc,
        "Draft:     Initiator \"\" | Vertical Head \"\" | MIS Coordinator \"\"\n"
        "Submit:    Initiator \"\" | Vertical Head Pending | MIS \"\"     (header Status = Pending)\n"
        "VH Approve: Initiator \"\" | Vertical Head Approved | MIS Pending\n"
        "MIS complete: ... MIS Approved                           (header Status = Approved)\n"
        "Rework:    pending role Rework; header Status = Rework; initiator edits and resubmits\n"
        "Reject:    pending role Rejected; header Status = Rejected (closed)",
    )
    para(
        doc,
        "applyNpdWorkflowAction (npdFormThunks → npdRequestGeneralInfoService) is the only mutation path. "
        "Approve/Rework/Reject from the form footer and from email action= query params share it. "
        "Actions are role-based and user-independent: any user assigned to the pending role may act "
        "(VH still brand-scoped; MIS any assigned MIS user). View mode never shows those buttons.",
    )
    heading(doc, "14.2 Workflow Status dialog (display only)", 2)
    para(
        doc,
        "NpdRequestListTable Workflow icon → useNpdWorkflowStatusDialog → buildNpdWorkflowStatusViewRows. "
        "Dialog: NpdWorkflowStatusDialog. Columns: Role, Name, Workflow Role, Status. Footer: Cancel only. "
        "No writes.",
    )
    bullets(
        doc,
        [
            "Role = WorkFlowJSON.Role (Initiator, Vertical Head, MIS Coordinator, …).",
            "Name = AuthorTitle if email matches author; else SharePoint siteUsers Title; else formatted email local-part.",
            "Workflow Role = “Initiator” or “Approver”.",
            "Initiator Status column shows Draft if the request is still Draft, otherwise Initiated (JSON Status is always empty for Initiator).",
            "Approver Status is Pending / Approved / Rework / Rejected / —.",
            "Every step with Status Pending is the current processing row: yellow background, italic orange role, orange Pending — matching the provided screenshot.",
        ],
    )
    para(
        doc,
        "Because the dialog reads the current WorkFlowJSON on the list row, it always reflects "
        "the latest saved approval step. After an approver acts, the next list refresh shows the new current user highlighted.",
    )

    heading(doc, "15. Permissions — how screens decide what you can do", 1)
    add_table(
        doc,
        ["Question", "Answered by"],
        [
            ["Which nav items?", "filterNavigationByRoles(selectResolvedAccess) — union of ApproversMaster roles + Admin group"],
            ["Can this URL open?", "ProtectedRoute + permissionService route map"],
            ["Which requests appear?", "canViewRequest / pending helpers in npdRequestListService"],
            ["Show Edit icon?", "canEditNpdListRow (own draft/rework or current pending VH/MIS step)"],
            ["Show Approve on form?", "mode=edit AND canActOnPendingNpdStep AND footer mode for that role — never in view"],
            ["Which brands on New Request?", "app.npdInitiatorBrands from ApproversMaster System=New Product Development, Role=Initiator"],
        ],
    )
    para(
        doc,
        "Multi-role users keep every role. Admin does not replace Initiator. If Leo is Vertical Head "
        "for ROCA only, he sees VH nav and ROCA-scoped lists — not Initiator New/Draft unless ApproversMaster "
        "also has him as Initiator.",
    )

    heading(doc, "16. Complete end-to-end flow", 1)
    heading(doc, "16.1 Login to first paint", 2)
    bullets(
        doc,
        [
            "SharePoint loads the web part. RocaNpdWebPart loads CSS and renders RocaNpd.",
            "PnP is bound to the SPFx context. Redux store wraps MainComponent.",
            "initializeApp stores the user and resolves roles from Admin group + ApproversMaster.",
            "AppShell renders. Side nav shows only modules for those roles.",
            "Default route is getDefaultRoute(access) (typically /npd/all or first allowed item).",
        ],
    )
    heading(doc, "16.2 Initiator creates a request", 2)
    bullets(
        doc,
        [
            "Opens /npd/new. Controller loads brand options (ApproversMaster), lookup options (NPD_Lookup), plants when Material Type is chosen.",
            "Fills General Information and Item Details (manual rows or Excel import).",
            "Save Draft writes NPD_Request + NPD_ItemDetails and rebuilds WorkFlowJSON with empty approver statuses. Appears on Draft / ReWork.",
            "Submit validates one field at a time, assigns NPD-YYYY-###, sets Pending, marks Vertical Head Pending, emails VH, shows line-item progress.",
            "Request leaves Draft list and appears on All + Pending (for VH).",
        ],
    )
    heading(doc, "16.3 Vertical Head processes it", 2)
    bullets(
        doc,
        [
            "Pending list shows requests whose WorkFlowJSON pending role is Vertical Head and Brand is in that user’s VH brands.",
            "Workflow icon shows Initiator = Initiated (green) and Vertical Head = Pending (highlighted).",
            "View is read-only. Edit opens approve/rework/reject.",
            "Approve writes VH Approved, sets MIS Coordinator Pending, header stays Pending, emails MIS (no action buttons).",
            "Rework/Reject collect comments into NPD_ApproverComments and update JSON + header Status.",
        ],
    )
    heading(doc, "16.4 MIS Coordinator processes it", 2)
    bullets(
        doc,
        [
            "Any user with MIS Coordinator in ApproversMaster sees the request when MIS is Pending (not limited to the JSON email).",
            "Workflow dialog now highlights MIS Coordinator as current.",
            "Edit allows Rework / Reject / Post to SAP (Post to SAP UI/SAP call still pending in TaskList).",
            "When the last approver completes, header Status becomes Approved and the row appears on Approved Requests.",
        ],
    )
    heading(doc, "16.5 How the pieces talk during that flow", 2)
    code(
        doc,
        "NpdDraftRework / NpdRequestDashboard\n"
        "  dispatch fetchNpdPendingList / fetchNpdDashboardList\n"
        "    npdRequestListService + npdRequestGeneralInfoService (NPD_Request + item counts)\n"
        "      npdRequestSlice.pendingItems / dashboardItems\n"
        "        NpdRequestListTable (Edit / Workflow / View)\n"
        "          Workflow → NpdWorkflowStatusDialog (WorkFlowJSON only)\n"
        "          View/Edit → /npd/new?id&mode&from\n"
        "            useNpdRequestFormController.hydrateNpdRequestForm\n"
        "              npdFormSlice\n"
        "                NpdGeneralInfoSection + NpdItemDetailsSection + Footer\n"
        "                  Save/Submit/Approve → npdFormThunks → generalInfo + item + workflow + email services\n"
        "                    SharePoint lists updated → user returns to `from` list → thunk refetch",
    )

    heading(doc, "17. Suggested code-review path", 1)
    bullets(
        doc,
        [
            "Start with src/Documents/TechnicalArchitecture.md §2, §6, §7, §12 and ProjectStandards R-NPD01–R-NPD21.",
            "Read Config.ts and Interface.ts so list/field names are not guessed.",
            "Read roleService.ts + permissionService.ts + ProtectedRoute — this is security.",
            "Read npdWorkflowJsonService.ts — this is the approval state machine.",
            "Read npdRequestGeneralInfoService.ts — persist + submit + apply action.",
            "Read NpdRequestListTable + NpdWorkflowStatusDialog — list UX including Workflow Status.",
            "Read useNpdRequestFormController.ts — form orchestration.",
            "Only then read individual section components and admin masters; they follow one repeated pattern.",
        ],
    )

    heading(doc, "18. Not implemented yet (do not expect these in this review)", 1)
    bullets(
        doc,
        [
            "Material Group module screens (placeholders only).",
            "Reports dashboard.",
            "Admin Material Master screen.",
            "MIS Other Details / SAP field auto-populate and live Post to SAP.",
            "Plant/Source filter by Brand Material Extension (T-0512).",
        ],
    )

    heading(doc, "19. Documents that govern this code", 1)
    add_table(
        doc,
        ["Document", "Use it for"],
        [
            ["TechnicalArchitecture.md", "System design, lists, workflow, role matrix — source of truth"],
            ["ProjectStandards.md", "Mandatory coding rules (wrappers, Redux path, NPD rules R-NPD*)"],
            ["TaskList.md", "What is done vs pending (T-0509 = Workflow Status dialog)"],
            ["Checklist.md", "Verification items (CHK-G08a = Workflow Status)"],
            ["This Word file", "Narrative of how the running code is wired for reviewers"],
        ],
    )

    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(
        "ROCA NPD  •  End-to-End Code Review Guide  •  Generated from the current codebase  •  18 September 2026"
    )
    set_run(run, size=9, color=MUTED)

    doc.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    build()
